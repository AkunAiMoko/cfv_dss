"""Convert the manuscript to DOCX with native tables and Word equations."""
from pathlib import Path
import re
import subprocess
import zipfile
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
source = (ROOT / 'article.md').read_text(encoding='utf-8')
output = ROOT / 'article.docx'
reference = ROOT / 'word_reference.docx'

template = Document()
section = template.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.75)
normal = template.styles['Normal']
normal.font.name, normal.font.size = 'Times New Roman', Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for name, size in [('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)]:
    style = template.styles[name]
    style.font.name, style.font.size = 'Times New Roman', Pt(size)
    style.font.bold = True
    style.paragraph_format.keep_with_next = True
template.save(reference)

# Markdown tables require separation from the preceding caption paragraph.
markdown = re.sub(r'([^\n])\n(?=\|)', r'\1\n\n', source)
# Rejoin consecutive table rows: only insert a blank before each table.
markdown = re.sub(r'(\|[^\n]*\|)\n\n(?=\|)', r'\1\n', markdown)
# Preserve equation numbers using portable TeX supported by the DOCX writer.
markdown = re.sub(r'\\tag\{(\d+)\}', lambda m: r'\qquad \text{(' + m.group(1) + ')}', markdown)
markdown = markdown.replace(r'|\hat{\theta}_j|', r'\left|\hat{\theta}_j\right|')
for relative in re.findall(r'!\[[^\]]*\]\(([^)]+)\)', source):
    assert (ROOT / relative).is_file(), f'Missing image: {relative}'
result = subprocess.run([
    'pandoc', '--from=markdown+tex_math_dollars+autolink_bare_uris-implicit_figures', '--to=docx',
    '--standalone', '--reference-doc=' + str(reference),
    '--resource-path=' + str(ROOT), '--output=' + str(output),
], input=markdown, text=True, encoding='utf-8', cwd=ROOT, capture_output=True, check=True)
if result.stderr:
    print(result.stderr)

doc = Document(output)
available = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
for shape in doc.inline_shapes:
    if shape.width > available:
        ratio = available / shape.width
        shape.height = int(shape.height * ratio)
        shape.width = available
for table_index, table in enumerate(doc.tables):
    table.style = 'Table Grid'
    table.autofit = False
    fractions = ([.17, .22, .33, .05, .23] if table_index == 0 else
                 [.25] + [.75 / (len(table.columns) - 1)] * (len(table.columns) - 1))
    for col, fraction in zip(table.columns, fractions):
        col.width = int(available * fraction)
    for row_index, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if row_index == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for cell, fraction in zip(row.cells, fractions):
            cell.width = int(available * fraction)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1
                for run in p.runs:
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True
in_references = False
for p in doc.paragraphs:
    if p.text.strip() == 'References':
        in_references = True
        continue
    if in_references and p.text.strip():
        p.paragraph_format.left_indent = Inches(.25)
        p.paragraph_format.first_line_indent = Inches(-.25)
        p.paragraph_format.space_after = Pt(5)
    if p._p.xpath('.//w:drawing'):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
    if p.text.startswith(('Figure ', 'Table ')):
        p.paragraph_format.keep_together = True
        if p.text.startswith('Table '):
            p.paragraph_format.keep_with_next = True
footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.save(output)

ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
with zipfile.ZipFile(output) as archive:
    xml = ET.fromstring(archive.read('word/document.xml'))
    media = [name for name in archive.namelist() if name.startswith('word/media/')]
    tables = xml.findall('.//w:tbl', ns)
    display_math = xml.findall('.//m:oMathPara', ns)
    equations = xml.findall('.//m:oMath', ns)
    assert len(media) == 7, f'Expected 7 embedded figures, got {len(media)}'
    assert len(tables) == 3, f'Expected 3 editable tables, got {len(tables)}'
    assert len(display_math) == 8, f'Expected 8 display equations, got {len(display_math)}'
    assert [len(t.findall('./w:tr', ns)) for t in tables] == [14, 9, 7]
    assert len(re.findall(r'https://doi.org/', archive.read('word/_rels/document.xml.rels').decode())) == 50
print(f'Created {output.name}: {len(media)} embedded figures, {len(tables)} editable tables, '
      f'{len(display_math)} display equations, {len(equations)} total Word math objects, 50 DOI links.')
