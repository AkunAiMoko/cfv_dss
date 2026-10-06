"""Convert the Indonesian guide and its PNG figures to an editable Word document."""
from pathlib import Path
import re
import subprocess
import zipfile
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
source = (ROOT / 'TAHAPAN_TEKNIS_RISET.md').read_text(encoding='utf-8')
output = ROOT / 'TAHAPAN_TEKNIS_RISET.docx'
images = re.findall(r'!\[[^\]]*\]\(([^)]+)\)', source)
assert all(path.endswith('.png') and (ROOT / path).is_file() for path in images)
assert '.svg' not in source and '```mermaid' not in source
# Explicit delimiters avoid ambiguity between absolute-value bars and Markdown syntax.
markdown = source.replace(r'|\widehat{\theta}_j|', r'\left|\widehat{\theta}_j\right|')
result = subprocess.run([
    'pandoc', '--from=markdown+tex_math_dollars+autolink_bare_uris-implicit_figures',
    '--to=docx', '--standalone',
    '--reference-doc=' + str(ROOT / 'manuscript' / 'word_reference.docx'),
    '--resource-path=' + str(ROOT), '--output=' + str(output),
], input=markdown, text=True, encoding='utf-8', cwd=ROOT, capture_output=True, check=True)
if result.stderr.strip():
    raise RuntimeError(result.stderr)

doc = Document(output)
section = doc.sections[0]
available = section.page_width - section.left_margin - section.right_margin
max_height = section.page_height - section.top_margin - section.bottom_margin - Inches(.8)
for shape in doc.inline_shapes:
    scale = min(1, available / shape.width, max_height / shape.height)
    shape.width, shape.height = int(shape.width * scale), int(shape.height * scale)

for table in doc.tables:
    table.style = 'Table Grid'
    table.autofit = False
    headers = [cell.text for cell in table.rows[0].cells]
    if headers[0] == 'Kode':
        fractions = [.17, .25, .5, .08]
    elif headers[0] == 'Berkas input':
        fractions = [.26, .24, .5]
    elif headers[0] == 'Halaman':
        fractions = [.22, .31, .47]
    elif headers[0] == 'Kebijakan':
        fractions = [.32, .68] if len(headers) == 2 else [.25, .4, .35]
    elif headers[0] == 'Model' and len(headers) == 5:
        fractions = [.36, .16, .16, .16, .16]
    else:
        fractions = [.29] + [.71 / (len(headers) - 1)] * (len(headers) - 1)
    for col, fraction in zip(table.columns, fractions):
        col.width = int(available * fraction)
    for i, row in enumerate(table.rows):
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if i == 0:
            props.append(OxmlElement('w:tblHeader'))
        for cell, fraction in zip(row.cells, fractions):
            cell.width = int(available * fraction)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True

for p in doc.paragraphs:
    if p._p.xpath('.//w:drawing'):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
    if p.text.startswith('Gambar alur'):
        p.paragraph_format.keep_together = True
        for run in p.runs:
            run.font.size = Pt(10)
    if p.style.name == 'Source Code':
        p.paragraph_format.keep_with_next = False
        for run in p.runs:
            run.font.size = Pt(8)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.save(output)

ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
with zipfile.ZipFile(output) as archive:
    assert archive.testzip() is None
    xml = ET.fromstring(archive.read('word/document.xml'))
    media = [path for path in archive.namelist() if path.startswith('word/media/')]
    tables = xml.findall('.//w:tbl', ns)
    equations = xml.findall('.//m:oMathPara', ns)
    expected_tables = len(re.findall(r'^\|[^\n]*\n\|[-:| ]+\|\s*$', source, re.M))
    assert len(media) == len(images) == 10
    assert all(path.endswith('.png') for path in media)
    assert len(tables) == expected_tables
    assert len(equations) == len(re.findall(r'^\$\$', source, re.M)) == 5
    text = ''.join(xml.itertext())
    assert 'Mermaid' not in text and '.svg' not in text
print(f'Created {output.name}: {len(media)} PNG images, {len(tables)} editable tables, '
      f'{len(equations)} display equations. Size: {output.stat().st_size:,} bytes.')
